import base64
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from datetime import UTC, datetime
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from threading import Barrier
from types import SimpleNamespace
from typing import cast
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import httpx
import pytest
from alembic import command as migrations
from alembic.config import Config
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import create_engine, inspect
from sqlalchemy.orm import Session

from job_apply_pro.api.routes.communications import get_communication_service
from job_apply_pro.config import get_settings
from job_apply_pro.domain.communications import (
    IntegrationProvider,
    MutationKind,
    MutationStatus,
    OutboundDraft,
)
from job_apply_pro.domain.knowledge import (
    CandidateDocument,
    CandidateDocumentVersionRecord,
    DocumentKind,
)
from job_apply_pro.domain.mail import (
    MAX_MAIL_ATTACHMENT_BYTES,
    MailAttachmentError,
)
from job_apply_pro.integrations.configuration import ProviderConnectionConfig
from job_apply_pro.integrations.provider_clients import GmailMessageProvider, OutlookMessageProvider
from job_apply_pro.main import create_app
from job_apply_pro.security.encryption import SensitiveDataCipher
from job_apply_pro.security.keys import StaticKeyProvider
from job_apply_pro.services.communications import CommunicationService
from job_apply_pro.services.mail_attachments import MailAttachmentResolver
from job_apply_pro.storage.communication_repository import CommunicationRepository
from job_apply_pro.storage.database import Base, get_session
from job_apply_pro.storage.models import OutboundDraftRow
from job_apply_pro.storage.repository_contracts import (
    CandidateKnowledgeRepositoryProtocol,
    WorkbenchRepositoryProtocol,
)
from test_mail_attachment_boundary import _command, _confirmation, _service


class _Knowledge:
    def __init__(self) -> None:
        self.documents: dict[str, CandidateDocument] = {}
        self.versions: dict[str, CandidateDocumentVersionRecord] = {}

    def get_document(self, document_id: str) -> CandidateDocument | None:
        return self.documents.get(document_id)

    def get_version_record(self, version_id: str) -> CandidateDocumentVersionRecord | None:
        return self.versions.get(version_id)


def _pdf_data() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(output)
    return output.getvalue()


class _Rig:
    def __init__(self, session: Session, tmp_path: Path, provider: IntegrationProvider) -> None:
        self.root = tmp_path / "documents"
        self.root.mkdir()
        self.knowledge = _Knowledge()
        self.workflow = SimpleNamespace(profile_id="profile-1")
        workflows = SimpleNamespace(
            get_snapshot=lambda key: self.workflow if key == "workflow-1" else None
        )
        self.cipher = SensitiveDataCipher(StaticKeyProvider(b"d" * 32))
        self.resolver = MailAttachmentResolver(
            cast(CandidateKnowledgeRepositoryProtocol, self.knowledge),
            cast(WorkbenchRepositoryProtocol, workflows),
            self.cipher,
            self.root,
        )
        _, self.repository, self.adapter, self.analysis_id = _service(session, provider)
        self.service = CommunicationService(
            self.repository,
            message_adapters={provider: self.adapter},
            attachment_resolver=self.resolver,
        )
        self.provider = provider

    def add(self, data: bytes | None = None, suffix: str = "pdf") -> str:
        if data is None:
            data = _pdf_data()
        version_id = f"version-{len(self.knowledge.versions) + 1}"
        document_id = f"document-{version_id}"
        now = datetime.now(UTC)
        self.knowledge.documents[document_id] = CandidateDocument(
            id=document_id,
            profile_id="profile-1",
            kind=DocumentKind.RESUME,
            display_name="Synthetic resume",
            variant_label="General",
            job_family_tags=[],
            is_primary=True,
            archived=False,
            created_at=now,
        )
        path = self.root / f"{version_id}.jap"
        path.write_text(
            self.cipher.encrypt_bytes(data, context=f"document:{version_id}:file"), encoding="ascii"
        )
        self.knowledge.versions[version_id] = CandidateDocumentVersionRecord(
            id=version_id,
            document_id=document_id,
            version=1,
            file_name=f"resume-{version_id}.{suffix}",
            media_type="application/pdf"
            if suffix == "pdf"
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            sha256=hashlib.sha256(data).hexdigest(),
            parser_version="synthetic-1",
            page_count=1,
            character_count=10,
            created_at=now,
            storage_path=str(path),
            encrypted_extraction="unused",
        )
        return version_id

    def draft(self, ids: list[str]) -> OutboundDraft:
        return self.service.create_draft(
            _command(self.provider, analysis_id=self.analysis_id).model_copy(
                update={"workflow_id": "workflow-1", "document_version_ids": ids}
            )
        )


class _Tokens:
    def access_token(self, provider: IntegrationProvider) -> str:
        return "synthetic-test-token"


@pytest.mark.parametrize("provider", [IntegrationProvider.GMAIL, IntegrationProvider.OUTLOOK])
def test_exact_verified_pdf_and_docx_bytes_are_serialized(
    session: Session, tmp_path: Path, provider: IntegrationProvider
) -> None:
    rig = _Rig(session, tmp_path, provider)
    pdf = _pdf_data()
    document = Document()
    document.add_paragraph("Synthetic cover letter")
    output = BytesIO()
    document.save(output)
    expected = [pdf, output.getvalue()]
    ids = [rig.add(pdf), rig.add(expected[1], "docx")]
    draft = rig.draft(ids)
    manifest = draft.attachment_manifest
    assert manifest is not None
    stored = session.get(OutboundDraftRow, draft.id)
    assert stored is not None
    assert "resume-version" not in stored.encrypted_payload
    assert rig.repository.get_draft(draft.id) == draft
    assert "provider_binding_fingerprint" not in draft.model_dump()
    requests: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        payload = json.loads(request.content)
        if provider is IntegrationProvider.GMAIL:
            assert request.url.path.endswith("/messages/send")
            raw = payload["raw"]
            mime = BytesParser(policy=policy.default).parsebytes(
                base64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4))
            )
            parts = list(mime.iter_attachments())
            assert [part.get_payload(decode=True) for part in parts] == expected
            assert [part.get_filename() for part in parts] == [
                item.file_name for item in manifest.attachments
            ]
            return httpx.Response(200, json={"id": "gmail-actual-message-id"})
        assert request.url.path.endswith("/me/sendMail")
        parts = payload["message"]["attachments"]
        assert [base64.b64decode(part["contentBytes"]) for part in parts] == expected
        assert all(part["@odata.type"] == "#microsoft.graph.fileAttachment" for part in parts)
        return httpx.Response(202)

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        adapter = (
            GmailMessageProvider(_Tokens(), client=client)
            if provider is IntegrationProvider.GMAIL
            else OutlookMessageProvider(_Tokens(), client=client)
        )
        service = CommunicationService(
            rig.repository, message_adapters={provider: adapter}, attachment_resolver=rig.resolver
        )
        result = service.send_draft(draft.id, _confirmation(draft))
    assert result.status is MutationStatus.ACCEPTED
    assert result.provider_resource_id == (
        "gmail-actual-message-id" if provider is IntegrationProvider.GMAIL else None
    )
    assert len(requests) == 1
    assert sorted(path.name for path in rig.root.iterdir()) == [f"{key}.jap" for key in ids]


@pytest.mark.parametrize(
    "mutation",
    [
        "owner",
        "archived",
        "kind",
        "hash",
        "name",
        "mime",
        "outside",
        "cipher",
        "context",
        "missing",
    ],
)
def test_preview_rejects_invalid_document_version(
    session: Session, tmp_path: Path, mutation: str
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    key = rig.add()
    version = rig.knowledge.versions[key]
    document = rig.knowledge.documents[version.document_id]
    changes: dict[str, dict[str, object]]
    if mutation in {"owner", "archived", "kind"}:
        changes = {
            "owner": {"profile_id": "profile-2"},
            "archived": {"archived": True},
            "kind": {"kind": DocumentKind.OTHER},
        }
        rig.knowledge.documents[document.id] = document.model_copy(update=changes[mutation])
    elif mutation in {"hash", "name", "mime", "outside"}:
        outside = tmp_path / "outside.jap"
        outside.write_text("encrypted fixture", encoding="ascii")
        changes = {
            "hash": {"sha256": "0" * 64},
            "name": {"file_name": "private\r\nBcc: secret.pdf"},
            "mime": {"media_type": "text/plain"},
            "outside": {"storage_path": str(outside)},
        }
        rig.knowledge.versions[key] = version.model_copy(update=changes[mutation])
    elif mutation == "missing":
        Path(version.storage_path).unlink()
    else:
        value = (
            "private-malformed-ciphertext"
            if mutation == "cipher"
            else rig.cipher.encrypt_bytes(b"%PDF-1.7", context="document:other:file")
        )
        Path(version.storage_path).write_text(value, encoding="ascii")
    with pytest.raises(MailAttachmentError) as error:
        rig.draft([key])
    assert "private" not in str(error.value)
    assert rig.service.list_drafts() == []
    assert rig.adapter.sent == []


@pytest.mark.parametrize("mutation", ["bytes", "hash", "name", "owner", "archive"])
def test_send_rechecks_manifest_and_authenticated_content(
    session: Session, tmp_path: Path, mutation: str
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    key = rig.add()
    draft = rig.draft([key])
    version = rig.knowledge.versions[key]
    if mutation in {"bytes", "hash"}:
        data = b"%PDF-1.7\nchanged after review"
        Path(version.storage_path).write_text(
            rig.cipher.encrypt_bytes(data, context=f"document:{key}:file"), encoding="ascii"
        )
        if mutation == "hash":
            rig.knowledge.versions[key] = version.model_copy(
                update={"sha256": hashlib.sha256(data).hexdigest()}
            )
    elif mutation == "name":
        rig.knowledge.versions[key] = version.model_copy(update={"file_name": "different.pdf"})
    else:
        document = rig.knowledge.documents[version.document_id]
        rig.knowledge.documents[document.id] = document.model_copy(
            update={"profile_id": "another"} if mutation == "owner" else {"archived": True}
        )
    with pytest.raises(MailAttachmentError):
        rig.service.send_draft(draft.id, _confirmation(draft))
    assert rig.adapter.sent == []
    assert rig.service.list_audits()[0].status is MutationStatus.FAILED


def test_missing_workflow_duplicate_count_and_aggregate_limits(
    session: Session, tmp_path: Path
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    key = rig.add()
    for workflow, ids in [
        (None, [key]),
        ("missing", [key]),
        ("workflow-1", [key, key]),
        ("workflow-1", [str(n) for n in range(5)]),
    ]:
        with pytest.raises(MailAttachmentError):
            rig.resolver.resolve(workflow, ids)
    data = _pdf_data()
    exact = rig.add(data + b"x" * (MAX_MAIL_ATTACHMENT_BYTES - len(data)))
    manifest, bundle = rig.resolver.resolve("workflow-1", [exact])
    assert manifest.attachments[0].size_bytes == MAX_MAIL_ATTACHMENT_BYTES
    assert len(bundle[0].data) == MAX_MAIL_ATTACHMENT_BYTES
    with pytest.raises(MailAttachmentError, match="2 MiB"):
        rig.resolver.resolve("workflow-1", [exact, key])
    oversized = rig.add(data + b"x" * MAX_MAIL_ATTACHMENT_BYTES)
    with pytest.raises(MailAttachmentError, match="2 MiB"):
        rig.resolver.resolve("workflow-1", [oversized])


@pytest.mark.parametrize("provider", [IntegrationProvider.GMAIL, IntegrationProvider.OUTLOOK])
def test_direct_provider_rechecks_resolved_bundle_before_token_access(
    session: Session, tmp_path: Path, provider: IntegrationProvider
) -> None:
    rig = _Rig(session, tmp_path, provider)
    key = rig.add()
    draft = rig.draft([key])
    _, bundle = rig.resolver.resolve("workflow-1", [key])
    bad = (replace(bundle[0], data=b"%PDF-1.7\nnot reviewed"),)

    class NoTokens:
        def access_token(self, provider: IntegrationProvider) -> str:
            raise AssertionError("must validate first")

    adapter = (
        GmailMessageProvider(NoTokens())
        if provider is IntegrationProvider.GMAIL
        else OutlookMessageProvider(NoTokens())
    )
    with pytest.raises(MailAttachmentError):
        adapter.send(draft, idempotency_key="bad-bundle-1", attachments=bad)
    assert "not reviewed" not in repr(bad)


@pytest.mark.parametrize("response_kind", ["timeout", "500", "redirect", "malformed", "oversized"])
def test_uncertain_send_is_persisted_and_never_resent(
    session: Session, tmp_path: Path, response_kind: str
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    draft = rig.draft([rig.add()])
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if response_kind == "timeout":
            raise httpx.ReadTimeout("private provider payload", request=request)
        if response_kind == "500":
            return httpx.Response(500, text="private provider payload")
        if response_kind == "redirect":
            return httpx.Response(302, headers={"location": "https://unexpected.example.test"})
        if response_kind == "oversized":
            return httpx.Response(200, content=b"x" * 70_000)
        return httpx.Response(200, json={"private": "no message id"})

    with httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=True) as client:
        service = CommunicationService(
            rig.repository,
            message_adapters={
                IntegrationProvider.GMAIL: GmailMessageProvider(_Tokens(), client=client)
            },
            attachment_resolver=rig.resolver,
        )
        confirmation = _confirmation(draft)
        result = service.send_draft(draft.id, confirmation)
        assert result.status is MutationStatus.UNCERTAIN
        assert result.provider_resource_id is None
        assert "private" not in result.model_dump_json()
        assert service.send_draft(draft.id, confirmation) == result
        with pytest.raises(ValueError, match="different send attempt"):
            service.send_draft(
                draft.id,
                confirmation.model_copy(update={"idempotency_key": "fresh-key-do-not-resend"}),
            )
    assert calls == 1


def test_idempotency_rejects_wrong_resource_fingerprint_and_account(
    session: Session, tmp_path: Path
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    first = rig.draft([rig.add()])
    confirmation = _confirmation(first)
    accepted = rig.service.send_draft(first.id, confirmation)
    for draft_id, command in [
        ("different-draft", confirmation),
        (first.id, confirmation.model_copy(update={"fingerprint": "b" * 64})),
    ]:
        with pytest.raises(ValueError, match="different send attempt"):
            rig.service.send_draft(draft_id, command)
    assert rig.service.send_draft(first.id, confirmation) == accepted
    second = rig.draft([next(iter(rig.knowledge.versions))])
    other_account = CommunicationService(
        rig.repository,
        message_adapters={IntegrationProvider.GMAIL: rig.adapter},
        attachment_resolver=rig.resolver,
        provider_configs={
            IntegrationProvider.GMAIL: ProviderConnectionConfig(
                provider=IntegrationProvider.GMAIL, credential_reference="oauth:different-account"
            )
        },
    )
    with pytest.raises(ValueError, match="provider account changed"):
        other_account.send_draft(
            second.id,
            _confirmation(second).model_copy(update={"idempotency_key": "other-account-key"}),
        )
    assert len(rig.adapter.sent) == 1


@pytest.mark.parametrize("has_attachments", [False, True])
def test_missing_provider_binding_never_sends_and_replays_failure(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, has_attachments: bool
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    draft = rig.draft([rig.add()] if has_attachments else [])
    malformed = draft.model_copy(update={"provider_binding_fingerprint": None})
    monkeypatch.setattr(rig.repository, "get_draft", lambda key: malformed)
    confirmation = _confirmation(draft)
    with pytest.raises(MailAttachmentError, match="provider-bound review"):
        rig.service.send_draft(draft.id, confirmation)
    assert rig.adapter.sent == []
    audits = rig.service.list_audits()
    assert len(audits) == 1
    assert audits[0].status is MutationStatus.FAILED
    assert rig.service.send_draft(draft.id, confirmation) == audits[0]
    assert rig.adapter.sent == []


def test_stored_metadata_validation_errors_do_not_expose_private_inputs(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)

    def malformed(key: str) -> None:
        raise ValueError("synthetic private stored filename and credential-like text")

    monkeypatch.setattr(rig.knowledge, "get_version_record", malformed)
    with pytest.raises(MailAttachmentError) as error:
        rig.draft(["version-1"])
    assert str(error.value) == "The selected document could not be verified"


def test_mail_send_reservation_is_atomic_and_blocks_legacy_effects(
    session: Session, tmp_path: Path
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    draft = rig.draft([rig.add()])
    first = rig.service._new_audit(
        kind=MutationKind.SEND_MESSAGE,
        provider=IntegrationProvider.GMAIL,
        resource_id=draft.id,
        command=_confirmation(draft),
    )
    reserved, claimed = rig.repository.claim_mail_send(first)
    assert claimed and reserved == first
    contender = first.model_copy(
        update={"id": "different-audit", "idempotency_key": "different-send-key"}
    )
    prior, claimed = rig.repository.claim_mail_send(contender)
    assert not claimed and prior == first
    with pytest.raises(ValueError, match="different send attempt"):
        rig.service.send_draft(
            draft.id, _confirmation(draft).model_copy(update={"idempotency_key": "third-send-key"})
        )
    assert rig.adapter.sent == []


def test_mail_claim_migration_roundtrip(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    url = f"sqlite:///{(tmp_path / 'mail-migration.db').as_posix()}"
    monkeypatch.setenv("JAP_DATABASE_URL", url)
    get_settings.cache_clear()
    config = Config("backend/alembic.ini")
    try:
        migrations.upgrade(config, "head")
        migrations.upgrade(config, "head")
        engine = create_engine(url)
        assert "mail_send_claims" in inspect(engine).get_table_names()
        migrations.downgrade(config, "20260814_0023")
        assert "mail_send_claims" not in inspect(engine).get_table_names()
        migrations.upgrade(config, "head")
        assert "mail_send_claims" in inspect(engine).get_table_names()
        engine.dispose()
    finally:
        get_settings.cache_clear()


@pytest.mark.parametrize(
    "kind",
    [
        DocumentKind.RESUME,
        DocumentKind.COVER_LETTER,
        DocumentKind.CERTIFICATION,
        DocumentKind.EDUCATION,
        DocumentKind.PORTFOLIO,
    ],
)
def test_all_requested_document_categories_can_be_reviewed(
    session: Session, tmp_path: Path, kind: DocumentKind
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    key = rig.add()
    document_id = rig.knowledge.versions[key].document_id
    rig.knowledge.documents[document_id] = rig.knowledge.documents[document_id].model_copy(
        update={"kind": kind}
    )
    draft = rig.draft([key])
    assert draft.attachment_manifest is not None
    assert draft.attachment_manifest.attachments[0].document_version_id == key


@pytest.mark.parametrize(
    "malformation",
    [
        "non_xml",
        "doctype",
        "utf16_doctype",
        "crc",
        "deflate",
        "unsafe_path",
        "macro_type",
        "duplicate_case",
    ],
)
def test_docx_container_and_required_xml_fail_closed(
    session: Session, tmp_path: Path, malformation: str
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    types = (
        b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        b'<Override PartName="/word/document.xml" ContentType="application/vnd.'
        b'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    )
    body = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'
    if malformation == "non_xml":
        body = b"not document XML"
    if malformation in {"doctype", "utf16_doctype"}:
        value = '<!DOCTYPE doc [<!ENTITY private "private secret">]>' + body.decode()
        body = value.encode("utf-16") if malformation == "utf16_doctype" else value.encode()
    if malformation == "macro_type":
        types = types.replace(
            b"openxmlformats-officedocument.wordprocessingml.document.main+xml",
            b"ms-word.document.macroEnabled.main+xml",
        )
    output = BytesIO()
    with ZipFile(
        output, "w", compression=ZIP_DEFLATED if malformation == "deflate" else ZIP_STORED
    ) as archive:
        archive.writestr("[Content_Types].xml", types)
        archive.writestr("word/document.xml", body)
        if malformation == "unsafe_path":
            archive.writestr("../private.bin", b"data")
        if malformation == "duplicate_case":
            archive.writestr("word/DOCUMENT.xml", body)
    data = output.getvalue()
    if malformation == "crc":
        data = data.replace(b"<w:body/>", b"<w:b0dy/>", 1)
    if malformation == "deflate":
        # Reserved DEFLATE block type must become a static failure, not a zlib exception.
        with ZipFile(BytesIO(data)) as archive:
            entry = archive.getinfo("word/document.xml")
            start = entry.header_offset + 30 + len(entry.filename.encode()) + len(entry.extra)
        data = data[:start] + b"\x07" + data[start + 1 :]
    with pytest.raises(MailAttachmentError, match="DOCX"):
        rig.draft([rig.add(data, "docx")])


def test_pdf_signature_only_input_is_rejected(session: Session, tmp_path: Path) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    with pytest.raises(MailAttachmentError, match="PDF"):
        rig.draft([rig.add(b"%PDF-1.7\nnot a PDF container")])


def test_other_mutation_cannot_overwrite_or_replay_mail_audit(
    session: Session, tmp_path: Path
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    draft = rig.draft([rig.add()])
    confirmation = _confirmation(draft)
    accepted = rig.service.send_draft(draft.id, confirmation)
    corrupted = accepted.model_copy(
        update={
            "kind": MutationKind.CREATE_CALENDAR_EVENT,
            "resource_id": "calendar-plan",
            "status": MutationStatus.FAILED,
        }
    )
    with pytest.raises(ValueError, match="different mutation"):
        rig.repository.add_audit(corrupted)
    with pytest.raises(ValueError, match="different mutation"):
        rig.service.execute_calendar_mutation("calendar-plan", confirmation)
    assert rig.repository.find_audit_by_idempotency(confirmation.idempotency_key) == accepted


def test_distinct_database_sessions_cannot_both_claim_one_send(tmp_path: Path) -> None:
    engine = create_engine(f"sqlite:///{(tmp_path / 'claims.db').as_posix()}")
    Base.metadata.create_all(engine)
    with Session(engine) as setup:
        rig = _Rig(setup, tmp_path, IntegrationProvider.GMAIL)
        draft = rig.draft([rig.add()])
        audit = rig.service._new_audit(
            kind=MutationKind.SEND_MESSAGE,
            provider=IntegrationProvider.GMAIL,
            resource_id=draft.id,
            command=_confirmation(draft),
        )
    barrier = Barrier(2)

    def claim(index: int) -> bool:
        with Session(engine) as session:
            repository = CommunicationRepository(
                session, SensitiveDataCipher(StaticKeyProvider(b"a" * 32))
            )
            contender = audit.model_copy(
                update={"id": f"claim-audit-{index}", "idempotency_key": f"claim-request-{index}"}
            )
            barrier.wait(timeout=5)
            _, owned = repository.claim_mail_send(contender)
            return owned

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(claim, [1, 2]))
    assert sum(results) == 1
    with Session(engine) as session:
        repository = CommunicationRepository(
            session, SensitiveDataCipher(StaticKeyProvider(b"a" * 32))
        )
        assert len(repository.list_audits()) == 1
    engine.dispose()


def test_api_reviews_exact_manifest_and_reports_provider_acceptance(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    rig = _Rig(session, tmp_path, IntegrationProvider.GMAIL)
    key = rig.add()
    command = _command(IntegrationProvider.GMAIL, analysis_id=rig.analysis_id).model_copy(
        update={"workflow_id": "workflow-1", "document_version_ids": [key]}
    )
    monkeypatch.setenv("JAP_API_TOKEN", "verified-mail-api-fixture")
    get_settings.cache_clear()
    app = create_app()
    app.dependency_overrides[get_communication_service] = lambda: rig.service
    app.dependency_overrides[get_session] = lambda: session
    headers = {"X-Job-Apply-Pro-Token": "verified-mail-api-fixture"}
    try:
        with TestClient(app) as client:
            response = client.post(
                "/api/v1/communications/drafts",
                headers=headers,
                json=command.model_dump(mode="json"),
            )
            assert response.status_code == 201
            draft = response.json()
            manifest = draft["attachment_manifest"]
            assert manifest["policy_version"] == "mail-attachments-v1"
            assert manifest["profile_id"] == "profile-1"
            assert manifest["attachments"][0]["sha256"] == rig.knowledge.versions[key].sha256
            assert "provider_binding_fingerprint" not in draft
            assert "storage_path" not in response.text
            read = client.get(f"/api/v1/communications/drafts/{draft['id']}", headers=headers)
            assert read.json() == draft
            result = client.post(
                f"/api/v1/communications/drafts/{draft['id']}/send",
                headers=headers,
                json={
                    "fingerprint": draft["fingerprint"],
                    "idempotency_key": "verified-mail-api-send",
                    "confirmed_by": "fixture",
                },
            )
            assert result.status_code == 200
            assert result.json()["status"] == "ACCEPTED"
            assert result.json()["provider_resource_id"] == "fixture-message-1"
    finally:
        app.dependency_overrides.clear()
        get_settings.cache_clear()
